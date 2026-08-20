#!/usr/bin/env python3
"""2026-2027 family school schedule — master + per-child + month-by-month.

Builds Family_School_Schedule_2026-2027.xlsx from the actual pacing plan
for Cody, Hannah, Caleb, Haley, Carson, and Heather. Plain values (no
cross-sheet formulas) so the workbook opens cleanly in Google Sheets and
every cell stays hand-editable.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Palette (matches the companion artifacts) ─────────────────────────────
INK        = "263129"
INK_SOFT   = "5B6B5E"
PAPER      = "EEF0E6"
LINE       = "D7DBC9"
ACCENT     = "3F6355"
ACCENT_TINT= "DEE8DD"
WHITE      = "FFFFFF"

STUDENTS = {
    "Cody":    {"grade": "Grade 7/8",  "color": "B0553F", "tint": "F4E6E1"},
    "Hannah":  {"grade": "Grade 7",    "color": "33698F", "tint": "E2ECF1"},
    "Caleb":   {"grade": "Grade 4/5",  "color": "93711F", "tint": "F1EAD7"},
    "Haley":   {"grade": "Grade 3",    "color": "6E4C96", "tint": "EBE4F2"},
    "Carson":  {"grade": "Grade 1",    "color": "3D7A62", "tint": "E1EEE8"},
    "Heather": {"grade": "Pre-K",      "color": "A14E76", "tint": "F3E4EC"},
}

DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri"]

MASTER_ROWS = {
    "Cody": [
        "Language Arts 7 · Math 7 · From Adam to Us",
        "Language Arts 7 · Math 7 · From Adam to Us",
        "Language Arts 7 · Math 7 · From Adam to Us",
        "Language Arts 7 · Math 7 · From Adam to Us",
        "Catch-up · Real Cool History for Kids podcast",
    ],
    "Hannah": [
        "Language Arts 7 · Math 6/7 · From Adam to Us",
        "Language Arts 7 · Math 6/7 · From Adam to Us",
        "Language Arts 7 · Math 6/7 · From Adam to Us",
        "Language Arts 7 · Math 6/7 · From Adam to Us",
        "Catch-up · Real Cool History for Kids podcast",
    ],
    "Caleb": [
        "Language Arts 4 · Math 4 · From Adam to Us",
        "Language Arts 4 · Math 4 · From Adam to Us",
        "Language Arts 4 · Math 4 · From Adam to Us",
        "Language Arts 4 · Math 4 · From Adam to Us",
        "Catch-up (+ floating buffer ~biweekly) · RCHK podcast",
    ],
    "Haley": [
        "Math 3 · Spelling · From Adam to Us · Devotional",
        "Math 3 · Reading · From Adam to Us · Devotional",
        "Math 3 · Spelling · From Adam to Us · Devotional",
        "Math 3 · Reading · From Adam to Us · Devotional",
        "Off — catch-up / leisure · RCHK podcast if joining",
    ],
    "Carson": [
        "Devotional · Tech apps",
        "Workbook · Devotional",
        "Devotional · Tech apps",
        "Workbook · Devotional",
        "Open / optional apps",
    ],
    "Heather": [
        "Devotional · Tech apps",
        "Workbook · Devotional",
        "Devotional · Tech apps",
        "Workbook · Devotional",
        "Open / optional apps",
    ],
}

# Per-child: (day, [subject bullets]) for the individual sheets
INDIVIDUAL_DAYS = {
    "Cody": [
        ("Mon", ["Language Arts 7", "Math 7", "From Adam to Us"]),
        ("Tue", ["Language Arts 7", "Math 7", "From Adam to Us"]),
        ("Wed", ["Language Arts 7", "Math 7", "From Adam to Us"]),
        ("Thu", ["Language Arts 7", "Math 7", "From Adam to Us"]),
        ("Fri", ["Catch-up / review", "Real Cool History for Kids podcast"]),
    ],
    "Hannah": [
        ("Mon", ["Language Arts 7", "Math 6 / Math 7", "From Adam to Us"]),
        ("Tue", ["Language Arts 7", "Math 6 / Math 7", "From Adam to Us"]),
        ("Wed", ["Language Arts 7", "Math 6 / Math 7", "From Adam to Us"]),
        ("Thu", ["Language Arts 7", "Math 6 / Math 7", "From Adam to Us"]),
        ("Fri", ["Catch-up / review", "Real Cool History for Kids podcast"]),
    ],
    "Caleb": [
        ("Mon", ["Language Arts 4", "Math 4", "From Adam to Us"]),
        ("Tue", ["Language Arts 4", "Math 4", "From Adam to Us"]),
        ("Wed", ["Language Arts 4", "Math 4", "From Adam to Us"]),
        ("Thu", ["Language Arts 4", "Math 4", "From Adam to Us"]),
        ("Fri", ["Catch-up / review", "Real Cool History for Kids podcast",
                 "Extra buffer day some weeks if behind"]),
    ],
    "Haley": [
        ("Mon", ["Math 3", "Spelling", "From Adam to Us", "Devotional (with Carson & Heather)"]),
        ("Tue", ["Math 3", "Reading", "From Adam to Us", "Devotional (with Carson & Heather)"]),
        ("Wed", ["Math 3", "Spelling", "From Adam to Us", "Devotional (with Carson & Heather)"]),
        ("Thu", ["Math 3", "Reading", "From Adam to Us", "Devotional (with Carson & Heather)"]),
        ("Fri", ["Off — catch-up or leisure", "Real Cool History for Kids podcast, if joining"]),
    ],
    "Carson": [
        ("Mon", ["Devotional (with Haley)", "Tech apps — ABC Reading / Todo Math"]),
        ("Tue", ["Workbook", "Devotional (with Haley)"]),
        ("Wed", ["Devotional (with Haley)", "Tech apps"]),
        ("Thu", ["Workbook", "Devotional (with Haley)"]),
        ("Fri", ["Open — optional tech apps, no set work"]),
    ],
    "Heather": [
        ("Mon", ["Devotional (with Haley)", "Tech apps — ABC Reading / Todo Math"]),
        ("Tue", ["Workbook", "Devotional (with Haley)"]),
        ("Wed", ["Devotional (with Haley)", "Tech apps"]),
        ("Thu", ["Workbook", "Devotional (with Haley)"]),
        ("Fri", ["Open — optional tech apps, no set work"]),
    ],
}

# Per-child: (subject, detail) for the "This Year" course list
COURSES = {
    "Cody": [
        ("Language Arts 7", "Books 4–10, ~5 wks/book, with Hannah"),
        ("Math 7", "Books 6–10, then review for the rest of the year"),
        ("From Adam to Us", "Science & History, group study Mon–Thu"),
    ],
    "Hannah": [
        ("Language Arts 7", "Books 4–10, ~5 wks/book, with Cody"),
        ("Math 6 → Math 7", "Math 6 books 6–10, then Math 7 books 1–3"),
        ("From Adam to Us", "Science & History, group study Mon–Thu"),
    ],
    "Caleb": [
        ("Language Arts 4", "Books 1–7, ~4.9 wks/book"),
        ("Math 4", "Books 1–7, ~4.9 wks/book"),
        ("From Adam to Us", "Science & History, group study Mon–Thu"),
    ],
    "Haley": [
        ("Math 3", "Books 2–10, ~5 wks/book"),
        ("Reading & Spelling", "2x/week each"),
        ("From Adam to Us", "Science & History, group study Mon–Thu"),
        ("Devotional", "Incl. Atlas devotional, with Carson & Heather"),
    ],
    "Carson": [
        ("Workbook", "Leisurely pace, 2x/week"),
        ("Devotional", "Incl. Atlas devotional, 4x/week with Haley"),
        ("Tech apps", "ABC Reading & Todo Math, 3x/week"),
    ],
    "Heather": [
        ("Workbook", "Leisurely pace, 2x/week"),
        ("Devotional", "Incl. Atlas devotional, 4x/week with Haley"),
        ("Tech apps", "ABC Reading & Todo Math, 3x/week"),
    ],
}

MONTHS = ["Sep '26", "Oct '26", "Nov '26", "Dec '26", "Break",
          "Jan '27", "Feb '27", "Mar '27", "Apr '27", "May '27"]

# (student, subject, [10 cells matching MONTHS, "" for the Break column])
MONTH_ROWS = [
    ("Cody", "Language Arts 7 (bks 4–10)",
        ["Bk 4", "Bk 5", "Bk 6", "Bk 6", "", "Bk 7", "Bk 8", "Bk 9", "Bk 9", "Bk 10"]),
    ("Cody", "Math 7 (bks 6–10, then review)",
        ["Bk 6", "Bk 7", "Bk 9", "Bk 9", "", "Bk 10", "Review", "Review", "Review", "Review"]),
    ("Hannah", "Language Arts 7 (bks 4–10)",
        ["Bk 4", "Bk 5", "Bk 6", "Bk 6", "", "Bk 7", "Bk 8", "Bk 9", "Bk 9", "Bk 10"]),
    ("Hannah", "Math 6 (bks 6–10) → Math 7 (bks 1–3)",
        ["M6 Bk6", "M6 Bk7", "M6 Bk9", "M6 Bk9", "", "M6 Bk10", "M7 Bk1", "M7 Bk2", "M7 Bk3", "M7 Bk3 done"]),
    ("Caleb", "Language Arts 4 (bks 1–7)",
        ["Bk 1", "Bk 2", "Bk 3", "Bk 3", "", "Bk 4", "Bk 5", "Bk 6", "Bk 6", "Bk 7 — clean finish"]),
    ("Caleb", "Math 4 (bks 1–7)",
        ["Bk 1", "Bk 2", "Bk 3", "Bk 3", "", "Bk 4", "Bk 5", "Bk 6", "Bk 6", "Bk 7 — clean finish"]),
    ("Haley", "Math 3 (bks 2–10)",
        ["Bk 2", "Bk 3", "Bk 4", "Bk 4", "", "Bk 5", "Bk 6", "Bk 7", "Bk 7", "Bk 8 (partial)"]),
    ("Haley", "Reading & Spelling (2x/wk each)",
        ["steady", "steady", "steady", "steady", "", "steady", "steady", "steady", "steady", "ongoing"]),
]

ASSUMPTIONS = [
    ("School year", "Tue, Sept 8, 2026 – Fri, May 28, 2027 (~34 instructional weeks)."),
    ("Maternity / new-baby break", "Last day Fri Dec 18, 2026 — resume Mon Jan 18, 2027."),
    ("Friday — Cody, Hannah, Caleb, Haley", "Catch-up/flex + Real Cool History for Kids podcast. New content runs Mon–Thu."),
    ("Friday — Carson & Heather", "Open / optional app time (ABC Reading, Todo Math) instead of catch-up — no new content to catch up on."),
    ("Math 7 & Math 6 pace (Cody & Hannah)", "Not specified in the original plan — assumed 4 weeks/book."),
    ("Language Arts 4 & Math 4 pace (Caleb)", "Re-paced to ~4.9 weeks/book so he finishes exactly book 7 at year end, not partway into book 8."),
    ("Caleb's catch-up day", "Floating, every 1–2 weeks — not pinned to a specific weekday. Using several may push his book 7 finish a bit past May 28."),
    ("Haley's weekly day off", "Same Friday catch-up day as everyone else, not a second day."),
    ("Haley's Math 3", "Not re-paced like Caleb's — books 9–10 will likely carry into next year at this pace."),
]


def fill(hex_):
    return PatternFill(start_color=hex_, end_color=hex_, fill_type="solid")

def font(size=10, bold=False, color=INK, italic=False):
    return Font(name="Calibri", size=size, bold=bold, color=color, italic=italic)

def align(h="left", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def thin_border(color=LINE):
    s = Side(style="thin", color=color)
    return Border(left=s, right=s, top=s, bottom=s)

def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def title_bar(ws, row, c1, c2, text, bg=ACCENT, fg=WHITE, size=16, height=32):
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    c = ws.cell(row, c1)
    c.value = text
    c.font = font(size=size, bold=True, color=fg)
    c.fill = fill(bg)
    c.alignment = align(h="center")
    ws.row_dimensions[row].height = height
    return c


def build_assumptions(wb):
    ws = wb.create_sheet("Assumptions")
    set_widths(ws, [34, 78])
    title_bar(ws, 1, 1, 2, "2026-2027 Family School Schedule — Assumptions")
    ws.row_dimensions[2].height = 6

    row = 3
    hdr = ["What", "Detail"]
    for col, text in enumerate(hdr, 1):
        c = ws.cell(row, col)
        c.value = text
        c.font = font(bold=True, color=WHITE)
        c.fill = fill(ACCENT)
        c.border = thin_border()
        c.alignment = align()
    row += 1

    for i, (what, detail) in enumerate(ASSUMPTIONS):
        bg = WHITE if i % 2 == 0 else PAPER
        c1 = ws.cell(row, 1, what)
        c1.font = font(bold=True, color=INK)
        c1.fill = fill(bg)
        c1.border = thin_border()
        c1.alignment = align(v="top", wrap=True)

        c2 = ws.cell(row, 2, detail)
        c2.font = font(color=INK)
        c2.fill = fill(bg)
        c2.border = thin_border()
        c2.alignment = align(v="top", wrap=True)

        ws.row_dimensions[row].height = 34
        row += 1

    ws.freeze_panes = "A4"
    return ws


def build_master(wb):
    ws = wb.create_sheet("Master Schedule")
    set_widths(ws, [16, 34, 34, 34, 34, 34])
    title_bar(ws, 1, 1, 6, "Master Schedule — Everyone, Mon–Fri", height=32)
    ws.row_dimensions[2].height = 6

    row = 3
    ws.cell(row, 1, "Student").font = font(bold=True, color=WHITE)
    ws.cell(row, 1).fill = fill(ACCENT)
    ws.cell(row, 1).border = thin_border()
    ws.cell(row, 1).alignment = align()
    for i, day in enumerate(DAYS):
        c = ws.cell(row, 2 + i, day)
        c.font = font(bold=True, color=WHITE)
        c.fill = fill(ACCENT)
        c.border = thin_border()
        c.alignment = align(h="center")
    row += 1

    for name, info in STUDENTS.items():
        name_c = ws.cell(row, 1, f"{name} · {info['grade']}")
        name_c.font = font(bold=True, color=WHITE)
        name_c.fill = fill(info["color"])
        name_c.border = thin_border()
        name_c.alignment = align(wrap=True)

        for i, text in enumerate(MASTER_ROWS[name]):
            c = ws.cell(row, 2 + i, text)
            c.font = font(italic=(i == 4), color=INK_SOFT if i == 4 else INK)
            c.fill = fill(info["tint"] if i == 4 else WHITE)
            c.border = thin_border()
            c.alignment = align(v="top", wrap=True)

        ws.row_dimensions[row].height = 48
        row += 1

    ws.freeze_panes = "B4"
    return ws


def build_student_sheet(wb, name):
    info = STUDENTS[name]
    ws = wb.create_sheet(name)
    set_widths(ws, [8, 60])

    title_bar(ws, 1, 1, 2, f"{name}  ·  {info['grade']}",
              bg=info["color"], height=32)
    ws.row_dimensions[2].height = 6

    row = 3
    ws.cell(row, 1, "Day").font = font(bold=True, color=WHITE)
    ws.cell(row, 1).fill = fill(info["color"])
    ws.cell(row, 1).border = thin_border()
    ws.cell(row, 1).alignment = align(h="center")
    ws.cell(row, 2, "Subjects").font = font(bold=True, color=WHITE)
    ws.cell(row, 2).fill = fill(info["color"])
    ws.cell(row, 2).border = thin_border()
    row += 1

    for day, subjects in INDIVIDUAL_DAYS[name]:
        is_fri = day == "Fri"
        d = ws.cell(row, 1, day)
        d.font = font(bold=True, color=INK_SOFT)
        d.fill = fill(info["tint"] if is_fri else WHITE)
        d.border = thin_border()
        d.alignment = align(h="center", v="top")

        s = ws.cell(row, 2, "\n".join(f"• {x}" for x in subjects))
        s.font = font(italic=is_fri, color=INK)
        s.fill = fill(info["tint"] if is_fri else WHITE)
        s.border = thin_border()
        s.alignment = align(v="top", wrap=True)

        ws.row_dimensions[row].height = 18 * len(subjects) + 10
        row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    lbl = ws.cell(row, 1, "This Year")
    lbl.font = font(bold=True, color=WHITE)
    lbl.fill = fill(ACCENT)
    lbl.alignment = align(h="center")
    row += 1

    ws.cell(row, 1, "Subject").font = font(bold=True, color=INK_SOFT)
    ws.cell(row, 2, "Detail").font = font(bold=True, color=INK_SOFT)
    for col in (1, 2):
        ws.cell(row, col).fill = fill(PAPER)
        ws.cell(row, col).border = thin_border()
    row += 1

    for i, (subject, detail) in enumerate(COURSES[name]):
        bg = WHITE if i % 2 == 0 else PAPER
        sc = ws.cell(row, 1, subject)
        sc.font = font(bold=True, color=INK)
        sc.fill = fill(bg)
        sc.border = thin_border()
        sc.alignment = align(v="top", wrap=True)

        dc = ws.cell(row, 2, detail)
        dc.font = font(color=INK_SOFT)
        dc.fill = fill(bg)
        dc.border = thin_border()
        dc.alignment = align(v="top", wrap=True)
        row += 1

    ws.sheet_properties.tabColor = info["color"]
    return ws


def build_month_by_month(wb):
    ws = wb.create_sheet("Month by Month")
    set_widths(ws, [10, 32] + [11] * len(MONTHS))
    n_cols = 2 + len(MONTHS)
    title_bar(ws, 1, 1, n_cols,
              "Month-by-Month Timeline — book/unit each subject is on by month end",
              height=32)
    ws.row_dimensions[2].height = 6

    row = 3
    hdr = ["Student", "Subject"] + MONTHS
    for col, text in enumerate(hdr, 1):
        c = ws.cell(row, col, text)
        c.font = font(bold=True, color=WHITE, size=9)
        c.fill = fill("A9A28C" if text == "Break" else ACCENT)
        c.border = thin_border()
        c.alignment = align(h="center", wrap=True)
    row += 1

    last_student = None
    for student, subject, cells in MONTH_ROWS:
        info = STUDENTS[student]
        new_block = student != last_student
        last_student = student

        name_c = ws.cell(row, 1, student if new_block else "")
        name_c.font = font(bold=True, color=WHITE)
        name_c.fill = fill(info["color"])
        name_c.border = thin_border()
        name_c.alignment = align(h="center")

        subj_c = ws.cell(row, 2, subject)
        subj_c.font = font(color=INK)
        subj_c.fill = fill(info["tint"])
        subj_c.border = thin_border()
        subj_c.alignment = align(v="center", wrap=True)

        for i, val in enumerate(cells):
            c = ws.cell(row, 3 + i, val)
            is_break = MONTHS[i] == "Break"
            c.font = font(size=9, color=INK_SOFT if (is_break or val in ("steady", "ongoing")) else INK,
                          bold=(val not in ("", "steady", "ongoing") and not is_break))
            c.fill = fill("E3E0D2" if is_break else WHITE)
            c.border = thin_border()
            c.alignment = align(h="center")

        ws.row_dimensions[row].height = 30
        row += 1

    ws.freeze_panes = "C4"
    return ws


def main():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    build_assumptions(wb)
    build_master(wb)
    for name in STUDENTS:
        build_student_sheet(wb, name)
    build_month_by_month(wb)

    wb["Assumptions"].sheet_properties.tabColor = ACCENT
    wb["Master Schedule"].sheet_properties.tabColor = ACCENT
    wb["Month by Month"].sheet_properties.tabColor = ACCENT

    out = "/home/user/curriculum-planning/Family_School_Schedule_2026-2027.xlsx"
    wb.save(out)
    print(f"Saved -> {out}")
    print(f"Sheets: {wb.sheetnames}")


if __name__ == "__main__":
    main()
